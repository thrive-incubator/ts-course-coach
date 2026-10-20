"""Course budget — fills the Thrive Academy Course Budget Worksheet from the user's answers."""

from __future__ import annotations

import io
import re
from pathlib import Path

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from pydantic import BaseModel, Field

router = APIRouter(prefix="/budget", tags=["budget"])

# The worksheet JB shared, kept verbatim so the download looks exactly like the
# spreadsheet the management team already uses. We only write into its cells.
_TEMPLATE_PATH = (
    Path(__file__).resolve().parent.parent / "templates" / "course_budget_template.xlsx"
)

FRINGE_RATE = 0.37
INFRASTRUCTURE_RATE = 0.30
ADMIN_SUPPORT_RATE = 0.15
FACULTY_ROWS = (26, 27, 28, 29)


class RevenueLine(BaseModel):
    learners: float = 0
    rate: float = 0
    notes: str = ""


class FacultyDirector(BaseModel):
    name: str = ""
    engagement_period: str = ""
    allocation: float = 0


class BudgetRequest(BaseModel):
    course_title: str = ""
    employer_sponsored: RevenueLine = Field(default_factory=RevenueLine)
    self_pay: RevenueLine = Field(default_factory=RevenueLine)
    scholarship: RevenueLine = Field(default_factory=RevenueLine)
    additional_revenue: float = 0
    additional_revenue_notes: str = ""
    faculty_directors: list[FacultyDirector] = Field(default_factory=list)
    faculty_sharing_notes: str = ""
    guest_speakers_count: float = 0
    guest_speaker_rate: float = 0
    guest_speaker_notes: str = ""
    additional_expenses: float = 0
    additional_expenses_notes: str = ""


def _slug(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return s[:40] or "course"


def build_workbook(req: BudgetRequest) -> bytes:
    wb = load_workbook(_TEMPLATE_PATH)
    ws = wb.active

    ws["A1"] = f"Course Title: {req.course_title}".rstrip()
    # The shared worksheet reads "$10,00" — a typo for $10,000.
    ws["A2"] = str(ws["A2"].value or "").replace("$10,00 ", "$10,000 ")

    # Revenue — quantity = learners, rate = tuition per learner.
    for row, line in ((5, req.employer_sponsored), (6, req.self_pay), (7, req.scholarship)):
        ws[f"B{row}"] = line.learners
        ws[f"C{row}"] = line.rate
        ws[f"D{row}"] = f"=B{row}*C{row}"
        ws[f"F{row}"] = line.notes or None
    ws["B8"] = "=SUM(B5:B7)"
    ws["D8"] = "=SUM(D5:D7)"
    ws["D9"] = req.additional_revenue
    ws["D9"].number_format = ws["D8"].number_format
    ws["F9"] = req.additional_revenue_notes or None

    # Faculty director distribution (up to four people) feeds salaries + fringe.
    for row, fd in zip(FACULTY_ROWS, req.faculty_directors[: len(FACULTY_ROWS)]):
        ws[f"A{row}"] = fd.name or None
        ws[f"B{row}"] = fd.engagement_period or None
        ws[f"C{row}"] = fd.allocation
        ws[f"D{row}"] = f"=C{row}*{FRINGE_RATE}"
    ws["D12"] = "=C30"
    ws["D13"] = "=D30"
    if req.faculty_sharing_notes:
        ws["F12"] = f"Indicate how teaching time is shared among faculty. {req.faculty_sharing_notes}"

    # Instructional costs.
    ws["B15"] = req.guest_speakers_count
    ws["C15"] = req.guest_speaker_rate
    ws["D15"] = "=B15*C15"
    ws["F15"] = req.guest_speaker_notes or None
    ws["D16"] = req.additional_expenses
    if req.additional_expenses_notes:
        ws["F16"] = req.additional_expenses_notes

    ws["D17"] = "=D12+D13+D15+D16"
    ws["D18"] = f"=D8*{INFRASTRUCTURE_RATE}"
    ws["D19"] = f"=D8*{ADMIN_SUPPORT_RATE}"
    ws["D22"] = "=IF(D8=0,0,D21/D8)"

    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.post("/xlsx")
async def download_budget(req: BudgetRequest) -> StreamingResponse:
    """Return the filled-in course budget worksheet as an .xlsx download."""
    data = build_workbook(req)
    filename = f"{_slug(req.course_title)}-course-budget.xlsx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
